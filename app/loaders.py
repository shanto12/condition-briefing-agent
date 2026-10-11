"""Document loaders for the RAG corpus and uploads: .md, .txt, .pdf, .xlsx, .csv -> LoadedDoc(meta, sections).

Pure functions: no network, no embedding, no PHI decision (app.rag runs the PHI gate on LoadedDoc.scan_text).
Metadata: .md front matter; other formats from manifest.json next to the file (keyed by filename), else derived
from the filename.
"""
from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

SUFFIXES = (".md", ".txt", ".pdf", ".xlsx", ".csv")
MIME_TYPES = {
    ".md": "text/markdown",
    ".txt": "text/plain",
    ".pdf": "application/pdf",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".csv": "text/csv",
}
MAX_BYTES = 10_000_000
MAX_PDF_PAGES = 50
MAX_SHEETS = 25
MAX_ROWS_PER_SHEET = 5_000
MAX_COLUMNS = 60
MAX_CHARS = 400_000
ROWS_PER_GROUP = 20
GROUP_CHARS = 2_500

_HEADING_MD = re.compile(r"^(#{1,3})\s+(.+?)\s*#*\s*$")
_HEADING_TXT = re.compile(r"^[A-Z][A-Z0-9 &/,'()\-]{3,79}$")
_BULLET = re.compile(r"^(\s*[-*\u2022]\s|\s*\d+[.)]\s)")


@dataclass
class LoadedDoc:
    meta: dict
    sections: list[tuple[str, str]]
    markdown: str | None = field(default=None, repr=False)

    @property
    def body(self) -> str:
        """Text for chunking: the original markdown for .md, otherwise one '## <label>' block per section."""
        if self.markdown is not None:
            return self.markdown
        return "\n\n".join(f"## {label}\n{text}" for label, text in self.sections)

    @property
    def scan_text(self) -> str:
        """Everything a PHI scan must see, metadata included."""
        return "\n".join(str(v) for v in self.meta.values()) + "\n" + self.body


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "document"


def load_path(path: Path) -> LoadedDoc:
    path = Path(path)
    if path.stat().st_size > MAX_BYTES:
        raise ValueError(f"{path.name}: file is larger than {MAX_BYTES // 1_000_000} MB")
    return load_bytes(path.name, path.read_bytes(), manifest(path.parent).get(path.name))


def load_bytes(filename: str, data: bytes, manifest_entry: dict | None = None) -> LoadedDoc:
    name = Path(filename or "").name
    suffix = Path(name).suffix.lower()
    if suffix not in SUFFIXES:
        raise ValueError(f"Unsupported file type '{suffix or name}'. Accepted: {', '.join(SUFFIXES)}.")
    if not data:
        raise ValueError(f"{name}: file is empty")
    if len(data) > MAX_BYTES:
        raise ValueError(f"{name}: file is larger than {MAX_BYTES // 1_000_000} MB")
    if suffix == ".md":
        doc = _load_markdown(name, data)
    elif suffix == ".txt":
        doc = _load_text(name, data, manifest_entry)
    elif suffix == ".pdf":
        doc = _load_pdf(name, data, manifest_entry)
    elif suffix == ".xlsx":
        doc = _load_xlsx(name, data, manifest_entry)
    else:
        doc = _load_csv(name, data, manifest_entry)
    if sum(len(t) for _, t in doc.sections) > MAX_CHARS:
        raise ValueError(f"{name}: more than {MAX_CHARS:,} characters of text")
    return doc


# ---------------- metadata ----------------

def manifest(folder: Path) -> dict:
    path = folder / "manifest.json"
    if not path.is_file():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    return data.get("files", data) if isinstance(data, dict) else {}


def _meta(name: str, fmt: str, found: dict | None, title_hint: str | None = None) -> dict:
    found = found or {}
    stem = Path(name).stem
    synthetic = found.get("synthetic", False)
    return {
        "doc_id": slug(str(found.get("doc_id") or stem)),
        "title": str(found.get("title") or title_hint or stem.replace("-", " ").replace("_", " ").title()),
        "doc_type": str(found.get("doc_type") or "document"),
        "owner": str(found.get("owner") or ""),
        "date": str(found.get("date") or ""),
        "synthetic": synthetic if isinstance(synthetic, bool) else str(synthetic).lower() == "true",
        "source_format": fmt,
    }


def _front_matter(text: str) -> tuple[dict, str]:
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not m:
        return {}, text
    meta = {}
    for line in m.group(1).splitlines():
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip()] = value.strip().strip('"')
    return meta, text[m.end():]


# ---------------- text formats ----------------

def _decode(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace").replace("\r\n", "\n")


def _load_markdown(name: str, data: bytes) -> LoadedDoc:
    front, body = _front_matter(_decode(data))
    sections, label, buf = [], "Notice", []
    for line in body.splitlines():
        m = _HEADING_MD.match(line)
        if m:
            if "".join(buf).strip():
                sections.append((label, "\n".join(buf).strip()))
            label, buf = m.group(2), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append((label, "\n".join(buf).strip()))
    return LoadedDoc(_meta(name, "md", front), sections, markdown=body)


def _load_text(name: str, data: bytes, entry: dict | None) -> LoadedDoc:
    front, body = _front_matter(_decode(data))
    sections, label, buf = [], "Notice", []
    for line in body.splitlines():
        if _HEADING_TXT.match(line.strip()) and not _BULLET.match(line):
            if "".join(buf).strip():
                sections.append((label, "\n".join(buf).strip()))
            label, buf = line.strip().title(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append((label, "\n".join(buf).strip()))
    return LoadedDoc(_meta(name, "txt", {**front, **(entry or {})}), sections)


# ---------------- PDF ----------------

def _reflow(lines: list[str]) -> list[str]:
    """Joins lines that a PDF wrapped mid-sentence, so sentence-level filters (injection) see whole sentences."""
    out: list[str] = []
    for line in (ln.rstrip() for ln in lines):
        if not line.strip():
            continue
        prev = out[-1] if out else ""
        wrapped = prev and not prev.endswith((".", ":", "!", "?")) and (len(prev) >= 60 or line[:1].islower())
        if wrapped and not _BULLET.match(line):
            out[-1] = f"{prev} {line.strip()}"
        else:
            out.append(line.strip())
    return out


def _load_pdf(name: str, data: bytes, entry: dict | None) -> LoadedDoc:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError(f"{name}: encrypted PDFs are not accepted")
        n_pages = len(reader.pages)
    except ValueError:
        raise
    except Exception as exc:  # pypdf raises several error types for malformed files
        raise ValueError(f"{name}: not a readable PDF ({type(exc).__name__})") from exc
    if n_pages > MAX_PDF_PAGES:
        raise ValueError(f"{name}: {n_pages} pages; the limit is {MAX_PDF_PAGES}")
    pages = [(page.extract_text() or "").splitlines() for page in reader.pages]
    repeated: set[str] = set()
    if len(pages) > 1:
        # Running headers and footers (page numbers normalized) appear on every page; keep them on page 1 only.
        norm = [{re.sub(r"\d+", "#", ln.strip()) for ln in p if ln.strip()} for p in pages]
        repeated = set.intersection(*norm)
    sections = []
    for i, lines in enumerate(pages, start=1):
        if i > 1:
            lines = [ln for ln in lines if re.sub(r"\d+", "#", ln.strip()) not in repeated]
        lines = _reflow(lines)
        if not lines:
            continue
        skip = repeated if i == 1 else set()
        heading = next((ln for ln in lines if re.sub(r"\d+", "#", ln) not in skip and len(ln) <= 80
                        and re.match(r"(\d+\.\s+)?[A-Z]", ln) and not ln.endswith(".")
                        and not ln.upper().startswith("SYNTHETIC")), "")
        sections.append((f"p{i} {heading}".strip(), "\n".join(lines)))
    info_title = None
    try:
        info_title = (reader.metadata or {}).get("/Title")
    except Exception:
        pass
    return LoadedDoc(_meta(name, "pdf", entry, info_title), sections)


# ---------------- spreadsheets ----------------

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ", "minutes")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return f"{value:.4f}".rstrip("0").rstrip(".")
    return re.sub(r"\s+", " ", str(value)).strip()


def _table_sections(name: str, sheet: str, rows) -> list[tuple[str, str]]:
    """First non-empty row is the header; each later row becomes 'Sheet <name> | <header>=<value>; ...'."""
    header: list[str] | None = None
    lines: list[tuple[int, str]] = []
    for n, row in enumerate(rows, start=1):
        if n > MAX_ROWS_PER_SHEET:
            raise ValueError(f"{name}: sheet '{sheet}' has more than {MAX_ROWS_PER_SHEET:,} rows")
        values = [_cell(v) for v in list(row)[:MAX_COLUMNS]]
        if not any(values):
            continue
        if header is None:
            header = [v or f"col{j + 1}" for j, v in enumerate(values)]
            continue
        pairs = [f"{header[j] if j < len(header) else f'col{j + 1}'}={v}" for j, v in enumerate(values) if v]
        lines.append((n, f"Sheet {sheet} | " + "; ".join(pairs)))
    if header is None:
        return []
    if not lines:
        return [(f"Sheet {sheet}", f"Sheet {sheet} | columns: {', '.join(header)}")]
    groups, current, size = [], [], 0
    for n, line in lines:
        if current and (len(current) >= ROWS_PER_GROUP or size + len(line) > GROUP_CHARS):
            groups.append(current)
            current, size = [], 0
        current.append((n, line))
        size += len(line) + 1
    groups.append(current)
    if len(groups) == 1:
        return [(f"Sheet {sheet}", "\n".join(line for _, line in groups[0]))]
    return [(f"Sheet {sheet} rows {g[0][0]}-{g[-1][0]}", "\n".join(line for _, line in g)) for g in groups]


def _load_xlsx(name: str, data: bytes, entry: dict | None) -> LoadedDoc:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # zipfile / openpyxl raise several error types for malformed files
        raise ValueError(f"{name}: not a readable .xlsx workbook ({type(exc).__name__})") from exc
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise ValueError(f"{name}: {len(wb.sheetnames)} sheets; the limit is {MAX_SHEETS}")
        sections = []
        for ws in wb.worksheets:
            sections += _table_sections(name, ws.title, ws.iter_rows(values_only=True))
        title = wb.properties.title
    finally:
        wb.close()
    return LoadedDoc(_meta(name, "xlsx", entry, title), sections)


def _load_csv(name: str, data: bytes, entry: dict | None) -> LoadedDoc:
    rows = csv.reader(io.StringIO(_decode(data)))
    return LoadedDoc(_meta(name, "csv", entry), _table_sections(name, Path(name).stem, rows))


if __name__ == "__main__":
    import sys

    folder = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "corpus")
    for p in sorted(x for x in folder.iterdir() if x.suffix.lower() in SUFFIXES):
        d = load_path(p)
        print(f"{d.meta['doc_id']:46} {d.meta['source_format']:5} {len(d.sections):3} sections  "
              f"synthetic={d.meta['synthetic']}  {[s[0] for s in d.sections][:3]}")
